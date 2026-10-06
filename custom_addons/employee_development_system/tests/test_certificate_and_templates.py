# -*- coding: utf-8 -*-
import base64
import io
import uuid
from openpyxl import Workbook
from odoo.tests.common import TransactionCase, tagged
from odoo.exceptions import ValidationError, UserError, AccessError


@tagged('post_install', '-at_install', 'eds', 'certificate', 'template_alignment')
class TestCertificateAndTemplates(TransactionCase):

    @classmethod
    def setUpClass(cls):
        super().setUpClass()
        # Create department
        cls.department = cls.env['hr.department'].create({
            'name': 'Talent Development Directorate',
        })

        # Create user & employee for participant
        cls.participant_user = cls.env['res.users'].with_context(no_reset_password=True, mail_create_nosubscribe=True).create({
            'name': 'Participant Employee',
            'login': 'participant_test_%s' % uuid.uuid4().hex[:8],
            'email': 'participant@bunnabank.et',
            'group_ids': [(6, 0, [cls.env.ref('base.group_user').id])],
        })
        cls.employee = cls.env['hr.employee'].create({
            'name': 'Abebe Bikila Test',
            'work_email': 'participant@bunnabank.et',
            'department_id': cls.department.id,
            'user_id': cls.participant_user.id,
            'identification_id': 'EMPTEST9901',
            'barcode': 'EMPTEST9901',
        })

        # Create EDS Officer user
        cls.officer_user = cls.env['res.users'].with_context(no_reset_password=True, mail_create_nosubscribe=True).create({
            'name': 'L&D Officer User',
            'login': 'officer_test_%s' % uuid.uuid4().hex[:8],
            'email': 'officer@bunnabank.et',
            'group_ids': [(6, 0, [
                cls.env.ref('base.group_user').id,
                cls.env.ref('employee_development_system.group_eds_officer').id,
            ])],
        })

        # Create Course
        cls.course = cls.env['eds.course'].create({
            'name': 'Credit Risk & Analysis Mastery',
            'code': 'CR-2026',
            'category': 'technical_compliance',
            'program_category': 'compliance',
            'target_audience': 'all_staff',
            'delivery_place': 'local',
            'duration_days': 3.0,
            'status': 'active',
        })

        # Create Session
        cls.session = cls.env['eds.session'].create({
            'name': 'CR-2026 Batch 01',
            'course_id': cls.course.id,
            'date_start': '2026-10-10 08:30:00',
            'date_end': '2026-10-12 17:30:00',
            'status': 'ongoing',
        })

        # Ensure Certificate Rule exists for technical_compliance
        cls.cert_rule = cls.env['eds.certificate.rule'].search([
            ('category', '=', 'technical_compliance')
        ], limit=1)
        if not cls.cert_rule:
            cls.cert_rule = cls.env['eds.certificate.rule'].create({
                'name': 'Technical Compliance Certification Rule',
                'category': 'technical_compliance',
                'require_attendance': True,
                'min_attendance_pct': 80.0,
                'require_level2_pass': True,
                'min_level2_score': 60.0,
                'active': True,
            })
        else:
            cls.cert_rule.write({
                'require_attendance': True,
                'min_attendance_pct': 80.0,
                'require_level2_pass': True,
                'min_level2_score': 60.0,
                'active': True,
            })

    def test_01_certificate_issuance_image_and_pdf(self):
        """Test issuing certificate generates PNG image, PDF, token, and delivery notification."""
        # Setup attendance 100%
        self.env['eds.session.attendance'].create({
            'session_id': self.session.id,
            'employee_id': self.employee.id,
            'attendance_date': '2026-10-10',
            'attended': True,
            'attendance_percentage': 100.0,
        })

        # Setup level 2 evaluation passing
        post_ass = self.env['eds.assessment'].create({
            'session_id': self.session.id,
            'employee_id': self.employee.id,
            'assessment_type': 'post',
            'score': 88.0,
        })
        self.env['eds.evaluation.level2'].create({
            'session_id': self.session.id,
            'employee_id': self.employee.id,
            'post_assessment_id': post_ass.id,
        })

        # Create certificate record
        cert = self.env['eds.certificate'].with_user(self.officer_user).create({
            'employee_id': self.employee.id,
            'session_id': self.session.id,
        })

        self.assertEqual(cert.state, 'pending')
        cert._compute_eligibility()
        self.assertTrue(cert.is_eligible, "Employee with 100% attendance and 88% L2 should be eligible")

        # Issue certificate as officer
        cert.action_issue()
        self.assertEqual(cert.state, 'issued')
        self.assertTrue(cert.verification_token, "Verification token must be generated")
        self.assertTrue(cert.verification_url, "Verification URL must be computed")
        self.assertIn(cert.verification_token, cert.verification_url)

        # Verify Certificate Image
        self.assertTrue(cert.certificate_image, "Certificate image binary must be populated")
        img_data = base64.b64decode(cert.certificate_image)
        # Verify PNG Magic Bytes
        self.assertTrue(img_data.startswith(b'\x89PNG\r\n\x1a\n'), "Generated image must be a valid PNG")
        self.assertGreater(len(img_data), 50000, "PNG image should have a realistic size (>50KB)")

        # Verify Certificate PDF
        self.assertTrue(cert.file, "Certificate PDF binary must be populated")
        self.assertTrue(len(cert.file) > 100, "PDF file must be non-empty")

        # Verify Notification was posted with attachments
        messages = cert.message_ids
        self.assertTrue(messages, "A notification message should be posted on the certificate record")
        attachments = messages[0].attachment_ids
        self.assertTrue(attachments, "Notification message should include attachments")

        # Test regeneration action
        old_image = cert.certificate_image
        cert.action_regenerate_assets()
        self.assertTrue(cert.certificate_image)
        self.assertTrue(cert.file)

    def test_02_certificate_ineligibility_blocks_issue(self):
        """Test participant who failed attendance cannot be issued a certificate."""
        emp2 = self.env['hr.employee'].create({
            'name': 'Ineligible Employee',
            'department_id': self.department.id,
        })
        # Low attendance: 50%
        self.env['eds.session.attendance'].create({
            'session_id': self.session.id,
            'employee_id': emp2.id,
            'attendance_date': '2026-10-10',
            'attended': False,
            'attendance_percentage': 50.0,
        })

        cert2 = self.env['eds.certificate'].with_user(self.officer_user).create({
            'employee_id': emp2.id,
            'session_id': self.session.id,
        })
        cert2._compute_eligibility()
        self.assertFalse(cert2.is_eligible)

        with self.assertRaises(ValidationError):
            cert2.action_issue()

    def test_03_portal_certificate_controller_verification(self):
        """Test public certificate verification controller logic."""
        # Create and issue a certificate
        self.env['eds.session.attendance'].create({
            'session_id': self.session.id,
            'employee_id': self.employee.id,
            'attended': True,
            'attendance_percentage': 100.0,
        })
        post_ass3 = self.env['eds.assessment'].create({
            'session_id': self.session.id,
            'employee_id': self.employee.id,
            'assessment_type': 'post',
            'score': 90.0,
        })
        self.env['eds.evaluation.level2'].create({
            'session_id': self.session.id,
            'employee_id': self.employee.id,
            'post_assessment_id': post_ass3.id,
        })
        cert = self.env['eds.certificate'].with_user(self.officer_user).create({
            'employee_id': self.employee.id,
            'session_id': self.session.id,
        })
        cert.action_issue()

        # Test searching by token
        found_by_token = self.env['eds.certificate'].sudo().search([
            ('verification_token', '=', cert.verification_token)
        ])
        self.assertEqual(found_by_token.id, cert.id)

        # Test searching by certificate code
        found_by_code = self.env['eds.certificate'].sudo().search([
            ('code', '=ilike', cert.code)
        ])
        self.assertEqual(found_by_code.id, cert.id)

        # Test void status
        cert.sudo().action_void()
        self.assertEqual(cert.state, 'void')
        self.assertFalse(cert.state == 'issued')

    def test_04_attendance_multi_slot_excel_import(self):
        """Test bulk attendance import wizard with multi-slot Excel format matching 1- Attendance.xlsx."""
        wb = Workbook()
        ws = wb.active
        ws.title = "Attendance"

        # Headers matching 1- Attendance.xlsx template
        ws.cell(1, 1, value="Bunna Bank S.C - Training Attendance Sheet")
        ws.cell(2, 1, value="Training Title: Credit Risk & Analysis Mastery")
        ws.cell(3, 1, value="Date: Oct 10-12, 2026")
        ws.cell(4, 1, value="")
        ws.cell(5, 1, value="No")
        ws.cell(5, 2, value="Participant Name")
        ws.cell(5, 3, value="Employee ID")
        ws.cell(5, 4, value="Morning")
        ws.cell(5, 5, value="Afternoon")
        ws.cell(5, 6, value="Morning")
        ws.cell(5, 7, value="Afternoon")

        # Row 6: Employee present for all slots
        ws.cell(6, 1, value=1)
        ws.cell(6, 2, value=self.employee.name)
        ws.cell(6, 3, value=self.employee.identification_id)
        ws.cell(6, 4, value="Present")
        ws.cell(6, 5, value="Present")
        ws.cell(6, 6, value="Present")
        ws.cell(6, 7, value="Present")

        buf = io.BytesIO()
        wb.save(buf)
        excel_data = base64.b64encode(buf.getvalue())

        wizard = self.env['eds.attendance.import'].create({
            'session_id': self.session.id,
            'csv_file': excel_data,
            'filename': '1- Attendance.xlsx',
        })
        wizard.action_import_attendance()

        self.assertIn("Excel Import Completed", wizard.result_log)
        att = self.env['eds.session.attendance'].search([
            ('session_id', '=', self.session.id),
            ('employee_id', '=', self.employee.id),
        ])
        self.assertTrue(att, "Attendance record should be created/updated by Excel import")
        self.assertTrue(att.attended)
        self.assertAlmostEqual(att.attendance_percentage, 100.0)

    def test_05_vendor_evaluation_matrix_threshold(self):
        """Test vendor evaluation matrix 40/70 technical threshold logic."""
        provider = self.env['eds.external.provider'].create({
            'name': 'FinTech Solutions Ltd',
            'provider_type': 'consultancy',
            'email': 'fintech@test.et',
        })

        rfp = self.env['eds.rfp'].create({
            'course_id': self.course.id,
            'scope_of_work': 'Provide Core Banking training modules.',
            'submission_deadline': '2026-11-30',
        })

        proposal = self.env['eds.vendor.proposal'].create({
            'rfp_id': rfp.id,
            'provider_id': provider.id,
        })
        self.assertEqual(proposal.state, 'registered')
        proposal.action_open_technical()
        self.assertEqual(proposal.state, 'technical_opened')

        # Create Technical Evaluation with score below 40 (e.g. 30)
        eval_fail = self.env['eds.vendor.evaluation'].create({
            'rfp_id': rfp.id,
            'proposal_id': proposal.id,
            'provider_id': provider.id,
            'evaluation_type': 'technical',
            'state': 'approved',
        })
        self.env['eds.vendor.evaluation.line'].create({
            'evaluation_id': eval_fail.id,
            'name': 'Technical Quality',
            'weight': 30.0,
            'score': 100.0,  # 30% contribution
        })
        eval_fail._compute_total_score()
        eval_fail._compute_threshold_met()

        proposal.technical_score_id = eval_fail
        self.assertFalse(proposal.technical_threshold_met, "Score of 30 is below technical threshold 40")

        # Attempting to open financial envelope must fail
        with self.assertRaises(UserError):
            proposal.action_open_financial()

        # Update evaluation to pass (e.g. score 50 >= 40)
        self.env['eds.vendor.evaluation.line'].create({
            'evaluation_id': eval_fail.id,
            'name': 'Experience & Track Record',
            'weight': 20.0,
            'score': 100.0,  # +20% -> total 50%
        })
        eval_fail._compute_total_score()
        eval_fail._compute_threshold_met()
        self.assertTrue(proposal.technical_threshold_met, "Score of 50 meets threshold of 40")

        # Now opening financial envelope succeeds
        proposal.action_open_financial()
        self.assertEqual(proposal.state, 'financial_opened')
