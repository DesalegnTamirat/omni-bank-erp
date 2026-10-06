# -*- coding: utf-8 -*-
import base64
import io
from openpyxl import Workbook
from odoo.tests.common import TransactionCase, tagged
from odoo.exceptions import ValidationError


@tagged('post_install', '-at_install', 'eds', 'template_alignment')
class TestTemplateAlignment(TransactionCase):

    @classmethod
    def setUpClass(cls):
        super().setUpClass()
        cls.department = cls.env['hr.department'].create({
            'name': 'Retail Banking Directorate',
        })
        cls.course = cls.env['eds.course'].create({
            'name': 'Anti-Money Laundering & KYC Compliance',
            'code': 'AML-001',
            'category': 'technical_compliance',
            'program_category': 'compliance',
            'target_audience': 'all_staff',
            'delivery_place': 'local',
            'delivery_method': 'internal',
            'duration_days': 2.0,
            'status': 'active',
        })
        cls.cycle = cls.env['eds.tna.cycle'].create({
            'name': 'TNA Cycle 2026/27',
            'year': '2026',
            'start_date': '2026-07-01',
            'submission_end_date': '2026-08-15',
            'approval_deadline': '2026-08-30',
            'state': 'approved',
        })

    def test_01_course_fields(self):
        """Test business template course fields."""
        self.assertEqual(self.course.program_category, 'compliance')
        self.assertEqual(self.course.target_audience, 'all_staff')
        self.assertEqual(self.course.delivery_place, 'local')

    def test_02_tna_entry_collective_fields(self):
        """Test TNA entry collective participant count and competency level."""
        tna = self.env['eds.tna.entry'].create({
            'cycle_id': self.cycle.id,
            'department_id': self.department.id,
            'course_id': self.course.id,
            'target_audience': 'below_mlm',
            'target_participant_count': 45,
            'competency_level': 'intermediate',
            'delivery_mode': 'workshop',
            'quarter': 'q1',
            'urgency': 'high',
            'justification': 'Mandatory regulatory AML refresher for all branch staff.',
        })
        self.assertEqual(tna.target_participant_count, 45)
        self.assertEqual(tna.competency_level, 'intermediate')
        self.assertEqual(tna.target_audience, 'below_mlm')

    def test_03_annual_plan_line_financial_formula(self):
        """Verify the exact financial formula matching Cal_BankLevel row 5:
        Total Provider Cost = Participants * Est Provider Cost
        Total Venue Cost = Participants * Duration * Est Venue Cost / Pax / Day
        Budget Allocated = Provider Cost + Venue Cost
        """
        plan = self.env['eds.annual.plan'].create({
            'name': 'Annual Plan 2026/27',
            'fiscal_year': '2026',
        })
        line = self.env['eds.annual.plan.line'].create({
            'plan_id': plan.id,
            'course_id': self.course.id,
            'program_name': 'AML & KYC Compliance Program',
            'program_category': 'compliance',
            'target_audience': 'all_staff',
            'scheduled_month': '2026-08',
            'duration_days': 3.0,
            'planned_participants': 30,
            'num_sessions': 2,
            'est_provider_cost_per_pax': 1200.0,
            'est_venue_cost_per_pax_day': 800.0,
        })

        # Quarter is computed from 2026-08 (August = Q1 in Bunna fiscal calendar)
        self.assertEqual(line.quarter, 'q1')
        # Total provider cost = 30 * 1200 = 36,000
        self.assertAlmostEqual(line.total_provider_cost, 36000.0)
        # Total venue cost = 30 * 3 * 800 = 72,000
        self.assertAlmostEqual(line.total_venue_cost, 72000.0)
        # Total budget = 36000 + 72000 = 108,000
        self.assertAlmostEqual(line.budget_allocated, 108000.0)

    def test_04_calendar_generation_populates_template_fields(self):
        """Verify action_generate_calendar maps target_audience, planned_participants, etc."""
        # Create approved TNA entry
        tna = self.env['eds.tna.entry'].create({
            'cycle_id': self.cycle.id,
            'department_id': self.department.id,
            'course_id': self.course.id,
            'target_audience': 'mlm',
            'target_participant_count': 25,
            'competency_level': 'advanced',
            'delivery_mode': 'classroom',
            'quarter': 'q2',
            'state': 'approved',
            'urgency': 'critical',
            'justification': 'Need for Q2',
        })

        plan = self.env['eds.annual.plan'].create({
            'name': 'Annual Plan Generated 2026/27',
            'fiscal_year': '2026',
            'state': 'draft',
        })
        plan.action_generate_calendar()

        # Check line created for the course
        lines = plan.line_ids.filtered(lambda l: l.course_id == self.course)
        self.assertTrue(lines, "Plan line should be generated from active course / approved TNA")
        line = lines[0]
        self.assertEqual(line.program_category, 'compliance')
        self.assertEqual(line.target_audience, 'all_staff')
        self.assertEqual(line.planned_participants, 25)

    def test_05_template_import_wizard(self):
        """Test multi-sheet excel template import wizard."""
        wb = Workbook()
        # Sheet 1: Lists
        ws_lists = wb.active
        ws_lists.title = "Lists"

        # Sheet 2: Curriculum
        ws_curr = wb.create_sheet(title="Curriculum")
        ws_curr.append(["Course Code", "Target Audience", "Focus Area", "Course Title", "Training Topics", "Days", "Delivery Place"])
        ws_curr.append(["TEST-CURR-01", "MLM", "Leadership", "Executive Management Mastery", "Module 1 Strategy", 3, "Local"])

        # Sheet 3: Provider_Register
        ws_prov = wb.create_sheet(title="Provider_Register")
        ws_prov.append(["Provider ID", "Provider Name", "Provider Type", "Specialisation Focus", "Contact Person", "Phone", "Email", "Indicative Rate/Session ETB", "Accreditation / Contract Ref"])
        ws_prov.append(["PROV-999", "Apex Consulting", "Consultancy", "Leadership & Finance", "Abebe Kebede", "+251911223344", "apex@consulting.et", 50000, "BUNNA-VEND-2026"])

        # Sheet 4: TNA
        ws_tna = wb.create_sheet(title="TNA")
        ws_tna.append(["Work Unit", "Training Topic / Course", "Target Group", "Estimated No of Participants", "Related Competency", "Competency Level", "Priority", "Preferred Delivery Method", "Preferred Quarter", "Remark"])
        ws_tna.append(["Retail Banking Directorate", "Customer Service Mastery", "Below MLM", 40, "Customer Service", "Intermediate", "High", "Workshop", "Q2 Oct-Dec", "Urgent"])

        # Save workbook to bytes
        buf = io.BytesIO()
        wb.save(buf)
        file_data = base64.b64encode(buf.getvalue())

        wizard = self.env['eds.template.import.wizard'].create({
            'file_name': 'EDS_LD_Templates_2026-27.xlsx',
            'file_data': file_data,
            'template_type': 'auto',
            'fiscal_year': '2026',
            'tna_cycle_id': self.cycle.id,
        })
        wizard.action_import()

        self.assertTrue(wizard.result_summary)
        # Check course created
        course = self.env['eds.course'].search([('code', '=', 'TEST-CURR-01')])
        self.assertTrue(course)
        self.assertEqual(course.name, "Executive Management Mastery")
        self.assertEqual(course.duration_days, 3.0)

        # Check provider created
        prov = self.env['eds.external.provider'].search([('name', '=', 'Apex Consulting')])
        self.assertTrue(prov)
        self.assertEqual(prov.email, 'apex@consulting.et')

        # Check TNA entry created
        tna = self.env['eds.tna.entry'].search([('proposed_program', '=', 'Customer Service Mastery'), ('cycle_id', '=', self.cycle.id)])
        self.assertTrue(tna)
        self.assertEqual(tna.target_participant_count, 40)
