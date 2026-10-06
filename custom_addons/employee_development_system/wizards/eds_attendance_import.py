# -*- coding: utf-8 -*-
from odoo import models, fields, _
from odoo.exceptions import ValidationError
import base64
import csv
import io
from datetime import datetime

try:
    import openpyxl
except ImportError:
    openpyxl = None


class EdsAttendanceImport(models.TransientModel):
    _name = 'eds.attendance.import'
    _description = 'Bulk Attendance Import Wizard'

    session_id = fields.Many2one('eds.session', string='Default Session', help='Optional if session code is in file')
    csv_file = fields.Binary(string='Attendance Sheet (CSV or Excel .xlsx)', required=True)
    filename = fields.Char(string='File Name')
    result_log = fields.Text(string='Import Summary', readonly=True)

    def action_download_template(self):
        """Return sample CSV template format."""
        sample_csv = (
            "session_ref,employee_badge,attendance_date,attended,notes\n"
            "SES/2026/0001,EMP001,2026-10-12,1,Present on time\n"
            "SES/2026/0001,EMP002,2026-10-12,0,Sick leave\n"
        )
        return {
            'type': 'ir.actions.act_url',
            'url': 'data:text/csv;charset=utf-8;base64,' + base64.b64encode(sample_csv.encode()).decode(),
            'target': 'self',
        }

    def action_import_attendance(self):
        self.ensure_one()
        if not self.csv_file:
            raise ValidationError(_("Please select an attendance CSV or Excel file to import."))

        data = base64.b64decode(self.csv_file)
        # Check if Excel XLSX (ZIP magic bytes PK\x03\x04)
        if data.startswith(b'PK\x03\x04') or (self.filename and self.filename.lower().endswith(('.xlsx', '.xlsm'))):
            return self._import_excel_attendance(data)
        else:
            return self._import_csv_attendance(data)

    def _import_excel_attendance(self, data):
        if not openpyxl:
            raise ValidationError(_("openpyxl Python library is required to import Excel attendance sheets."))

        try:
            wb = openpyxl.load_workbook(io.BytesIO(data), data_only=True)
        except Exception as e:
            raise ValidationError(_("Failed to open Excel workbook: %s") % str(e))

        AttendanceModel = self.env['eds.session.attendance']
        EmployeeModel = self.env['hr.employee']
        SessionModel = self.env['eds.session']

        success_count = 0
        skipped_count = 0
        errors = []

        for sname in wb.sheetnames:
            ws = wb[sname]
            if ws.max_row < 6:
                continue

            # Detect Session from Title in rows 1-6
            session = self.session_id
            if not session:
                for r in range(1, min(7, ws.max_row + 1)):
                    for c in range(1, min(5, ws.max_column + 1)):
                        val = str(ws.cell(r, c).value or '').strip()
                        if 'training title' in val.lower() or 'program' in val.lower():
                            parts = val.split(':', 1)
                            title_candidate = parts[1].strip() if len(parts) > 1 else val
                            # Search session by course name or session name
                            found_s = SessionModel.search([
                                '|',
                                ('name', '=ilike', title_candidate),
                                ('course_id.name', '=ilike', title_candidate),
                            ], limit=1)
                            if found_s:
                                session = found_s
                                break
                    if session:
                        break

            if not session:
                session = self.session_id

            # Find participant header row
            header_row = None
            col_name = None
            col_id = None
            slot_cols = []

            for r in range(1, min(12, ws.max_row + 1)):
                row_texts = {c: str(ws.cell(r, c).value or '').strip().lower() for c in range(1, ws.max_column + 1)}
                for c, text in row_texts.items():
                    if any(k in text for k in ('badge', 'emp id', 'employee id', 'staff id')) or text in ('id', 'code'):
                        col_id = c
                        header_row = r
                    elif any(k in text for k in ('participant', 'trainee', 'staff name', 'full name', 'employee name')) or ('name' in text and not col_name):
                        col_name = c
                        header_row = r
                    elif 'employee' in text and not col_id and not col_name:
                        col_name = c
                        header_row = r
                if header_row and (col_name or col_id):
                    break

            if not header_row:
                continue

            # Look for morning/afternoon slots in rows header_row, header_row+1, header_row+2
            for check_r in range(header_row, min(header_row + 3, ws.max_row + 1)):
                for c in range(1, ws.max_column + 1):
                    val = str(ws.cell(check_r, c).value or '').strip().lower()
                    if val in ('morning', 'afternoon', 'signature', 'present') and c not in slot_cols:
                        slot_cols.append(c)

            if not slot_cols:
                # Default slot cols to any columns to the right of Name/ID
                min_slot_c = max(col_name or 2, col_id or 2) + 1
                slot_cols = list(range(min_slot_c, min(min_slot_c + 6, ws.max_column + 1)))

            # Read participants
            for r in range(header_row + 1, ws.max_row + 1):
                p_name = str(ws.cell(r, col_name or 2).value or '').strip() if col_name else ''
                p_id = str(ws.cell(r, col_id or 3).value or '').strip() if col_id else ''

                if not p_name and not p_id:
                    continue
                # Skip sub-headers or footers
                if any(k in p_name.lower() for k in ('name of', 'morning', 'afternoon', 'signature', 'facilitated', 'prepared', 'approved')):
                    continue

                # Match employee
                emp = False
                if p_id:
                    emp = EmployeeModel.search([
                        '|',
                        ('identification_id', '=ilike', p_id),
                        ('barcode', '=ilike', p_id)
                    ], limit=1)
                if not emp and p_name:
                    emp = EmployeeModel.search([('name', '=ilike', p_name)], limit=1)

                if not emp:
                    errors.append(_("Sheet '%s' Row %d: Employee '%s' (%s) not found.") % (sname, r, p_name, p_id))
                    skipped_count += 1
                    continue

                if not session:
                    errors.append(_("Sheet '%s': No session mapped for participant %s.") % (sname, emp.name))
                    skipped_count += 1
                    continue

                # Calculate attended slots
                present_slots = 0
                total_slots = len(slot_cols) if slot_cols else 1
                for sc in slot_cols:
                    s_val = ws.cell(r, sc).value
                    # Count non-empty signatures or presence marks
                    if s_val is not None and str(s_val).strip() not in ('', '0', 'absent', 'a', 'false'):
                        present_slots += 1

                # If no slot markers filled in, treat employee row as present by default (full attendance)
                att_pct = (present_slots / total_slots * 100.0) if present_slots > 0 else 100.0

                existing = AttendanceModel.search([
                    ('session_id', '=', session.id),
                    ('employee_id', '=', emp.id)
                ], limit=1)

                if existing:
                    existing.write({
                        'attended': True,
                        'attendance_percentage': att_pct,
                        'attendance_date': fields.Date.today(),
                    })
                else:
                    AttendanceModel.create({
                        'session_id': session.id,
                        'employee_id': emp.id,
                        'attended': True,
                        'attendance_percentage': att_pct,
                        'attendance_date': fields.Date.today(),
                    })
                success_count += 1

        summary = _("Excel Import Completed: %d participant attendance records processed, %d skipped.") % (success_count, skipped_count)
        if errors:
            summary += "\n\n" + _("Details:\n") + "\n".join(errors[:25])
        self.result_log = summary

        return {
            'type': 'ir.actions.act_window',
            'res_model': 'eds.attendance.import',
            'res_id': self.id,
            'view_mode': 'form',
            'target': 'new',
        }

    def _import_csv_attendance(self, data):
        try:
            file_input = io.StringIO(data.decode('utf-8-sig'))
            reader = csv.DictReader(file_input)
        except Exception as e:
            raise ValidationError(_("Failed to read CSV file: %s") % str(e))

        success_count = 0
        skipped_count = 0
        errors = []
        AttendanceModel = self.env['eds.session.attendance']
        EmployeeModel = self.env['hr.employee']
        SessionModel = self.env['eds.session']

        rows = list(reader)
        session_refs = set(r.get('session_ref', '').strip() for r in rows if r.get('session_ref', '').strip())
        badges = set(r.get('employee_badge', '').strip() for r in rows if r.get('employee_badge', '').strip())

        # Pre-fetch sessions
        session_map = {}
        if session_refs:
            for s in SessionModel.search([('name', 'in', list(session_refs))]):
                session_map[s.name] = s
        if self.session_id:
            session_map[self.session_id.name] = self.session_id

        # Pre-fetch employees
        employee_map = {}
        if badges:
            emps = EmployeeModel.search([
                '|', '|',
                ('identification_id', 'in', list(badges)),
                ('barcode', 'in', list(badges)),
                ('name', 'in', list(badges)),
            ])
            for emp in emps:
                if emp.identification_id:
                    employee_map[emp.identification_id] = emp
                if emp.barcode:
                    employee_map[emp.barcode] = emp
                if emp.name:
                    employee_map[emp.name] = emp

        session_ids = [s.id for s in session_map.values()]
        existing_att_map = {}
        if session_ids:
            for att in AttendanceModel.search([('session_id', 'in', session_ids)]):
                existing_att_map[(att.session_id.id, att.employee_id.id)] = att

        new_vals_list = []
        for line_no, row in enumerate(rows, start=2):
            session_ref = row.get('session_ref', '').strip()
            badge = row.get('employee_badge', '').strip()
            date_str = row.get('attendance_date', '').strip()
            attended_str = row.get('attended', '').strip().lower()

            session = session_map.get(session_ref) or self.session_id
            if not session:
                errors.append(_("Row %d: Session '%s' not found.") % (line_no, session_ref))
                skipped_count += 1
                continue

            employee = employee_map.get(badge)
            if not employee:
                errors.append(_("Row %d: Employee '%s' not found.") % (line_no, badge))
                skipped_count += 1
                continue

            att_date = fields.Date.today()
            if date_str:
                try:
                    att_date = datetime.strptime(date_str, '%Y-%m-%d').date()
                except ValueError:
                    try:
                        att_date = datetime.strptime(date_str, '%d/%m/%Y').date()
                    except ValueError:
                        pass

            is_attended = attended_str in ('1', 'true', 'yes', 'present', 'y')

            existing_att = existing_att_map.get((session.id, employee.id))
            if existing_att:
                existing_att.write({
                    'attended': is_attended,
                    'attendance_date': att_date,
                    'attendance_percentage': 100.0 if is_attended else 0.0,
                })
                success_count += 1
            else:
                new_vals_list.append({
                    'session_id': session.id,
                    'employee_id': employee.id,
                    'attendance_date': att_date,
                    'attended': is_attended,
                    'attendance_percentage': 100.0 if is_attended else 0.0,
                })
                success_count += 1

        if new_vals_list:
            AttendanceModel.create(new_vals_list)

        summary = _("CSV Import Completed: %d attendance records processed, %d skipped.") % (success_count, skipped_count)
        if errors:
            summary += "\n\n" + _("Details:\n") + "\n".join(errors[:20])
        self.result_log = summary

        return {
            'type': 'ir.actions.act_window',
            'res_model': 'eds.attendance.import',
            'res_id': self.id,
            'view_mode': 'form',
            'target': 'new',
        }
