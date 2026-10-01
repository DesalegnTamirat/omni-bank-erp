# -*- coding: utf-8 -*-
import base64
import io
import logging

from odoo import api, fields, models, _
from odoo.exceptions import UserError

_logger = logging.getLogger(__name__)

try:
    from openpyxl import load_workbook, Workbook
    from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
except ImportError:  # pragma: no cover
    load_workbook = None
    Workbook = None


class PbmsExcelImportWizard(models.TransientModel):
    """Bulk-import and template generator for all 7 PBMS planning categories:
    Deposit, Customer Base, FX, Digital Banking, General Expense, Manpower, Fixed Asset.
    """
    _name = "pbms.excel.import.wizard"
    _description = "Import Plan & Budget Excel"

    plan_id = fields.Many2one("pbms.planning.category", string="Target Plan Card", readonly=True)
    cycle_id = fields.Many2one("pbms.planning.cycle", string="Planning Cycle", required=True)
    org_unit_id = fields.Many2one(
        "operating.unit", string="Operating Unit", required=True,
        default=lambda self: self.env.user.default_operating_unit_id,
    )
    format_type = fields.Selection(
        [
            ("deposit", "Deposit Mobilization"),
            ("customer_base", "Customer Base"),
            ("fx", "FX Mobilization"),
            ("digital_banking", "Digital Banking"),
            ("general_expense", "General Expense"),
            ("manpower", "Manpower Planning"),
            ("fixed_asset", "Fixed Asset Requirement"),
        ],
        string="Category",
        required=True,
        default="deposit",
    )
    import_mode = fields.Selection(
        [
            ("replace", "Replace existing lines for this category"),
            ("append", "Add to existing lines"),
        ],
        string="Import Mode",
        default="replace",
        required=True,
    )
    file_data = fields.Binary(string="Excel File (.xlsx)")
    file_name = fields.Char(string="File Name")

    @api.model
    def default_get(self, fields_list):
        res = super().default_get(fields_list)
        active_model = self.env.context.get("active_model")
        active_id = self.env.context.get("active_id")
        if active_model == "pbms.planning.category" and active_id:
            plan = self.env["pbms.planning.category"].browse(active_id)
            if plan.exists():
                res["plan_id"] = plan.id
                res["cycle_id"] = plan.cycle_id.id
                res["org_unit_id"] = plan.org_unit_id.id
                if plan.category:
                    res["format_type"] = plan.category
        elif self.env.context.get("default_category"):
            res["format_type"] = self.env.context.get("default_category")
        elif not res.get("cycle_id"):
            cycle = self.env["pbms.planning.cycle"].search([("state", "=", "open")], limit=1)
            if cycle:
                res["cycle_id"] = cycle.id
        return res

    def _get_target_plan(self):
        self.ensure_one()
        if self.plan_id:
            return self.plan_id
        Plan = self.env["pbms.planning.category"]
        plan = Plan.search([
            ("cycle_id", "=", self.cycle_id.id),
            ("org_unit_id", "=", self.org_unit_id.id),
            ("category", "=", self.format_type),
        ], limit=1)
        if not plan:
            plan = Plan.create({
                "cycle_id": self.cycle_id.id,
                "org_unit_id": self.org_unit_id.id,
                "category": self.format_type,
            })
        return plan

    def action_download_template(self):
        """Generate and return pre-formatted Excel template with sample master data for the chosen category."""
        self.ensure_one()
        if Workbook is None:
            raise UserError(_("The 'openpyxl' Python library is required. Please install it on the server."))

        wb = Workbook()
        ws = wb.active
        ws.title = "Template"

        # Theme styling
        fill_header = PatternFill(start_color="541718", end_color="541718", fill_type="solid")
        font_header = Font(name="Segoe UI", size=11, bold=True, color="FFFFFF")
        align_center = Alignment(horizontal="center", vertical="center")

        cat = self.format_type
        if cat in ("deposit", "customer_base"):
            headers = ["Product Code", "Product Name", "Opening Balance",
                       "Jul", "Aug", "Sep", "Oct", "Nov", "Dec", "Jan", "Feb", "Mar", "Apr", "May", "Jun"]
            ws.append(headers)
            dep_types = self.env["pbms.deposit.type"].search([("active", "=", True)], order="sequence, id")
            for dt in dep_types:
                ws.append([dt.code or "", dt.name or "", 0.0, 0.0, 0.0, 0.0, 0.0, 0.0, 0.0, 0.0, 0.0, 0.0, 0.0, 0.0, 0.0])
        elif cat == "fx":
            headers = ["Source Code", "FX Source Name",
                       "Jul", "Aug", "Sep", "Oct", "Nov", "Dec", "Jan", "Feb", "Mar", "Apr", "May", "Jun"]
            ws.append(headers)
            fx_types = self.env["pbms.fx.source.type"].search([("active", "=", True)])
            for fx in fx_types:
                ws.append([fx.code or "", fx.name or "", 0.0, 0.0, 0.0, 0.0, 0.0, 0.0, 0.0, 0.0, 0.0, 0.0, 0.0, 0.0, 0.0])
        elif cat == "digital_banking":
            headers = ["Channel Code", "Channel Name",
                       "Jul", "Aug", "Sep", "Oct", "Nov", "Dec", "Jan", "Feb", "Mar", "Apr", "May", "Jun"]
            ws.append(headers)
            channels = self.env["pbms.digital.channel"].search([("active", "=", True)])
            for ch in channels:
                ws.append([ch.code or "", ch.name or "", 0.0, 0.0, 0.0, 0.0, 0.0, 0.0, 0.0, 0.0, 0.0, 0.0, 0.0, 0.0, 0.0])
        elif cat == "general_expense":
            headers = ["Account Code", "Expense Account Name",
                       "Jul", "Aug", "Sep", "Oct", "Nov", "Dec", "Jan", "Feb", "Mar", "Apr", "May", "Jun"]
            ws.append(headers)
            accounts = self.env["pbms.expense.account"].search([("active", "=", True)], limit=20)
            for acc in accounts:
                ws.append([acc.code or "", acc.name or "", 0.0, 0.0, 0.0, 0.0, 0.0, 0.0, 0.0, 0.0, 0.0, 0.0, 0.0, 0.0, 0.0])
        elif cat == "manpower":
            headers = ["Position Type", "Job Title", "Authorized Baseline",
                       "Jul", "Aug", "Sep", "Oct", "Nov", "Dec", "Jan", "Feb", "Mar", "Apr", "May", "Jun",
                       "Monthly Compensation", "Justification Reason"]
            ws.append(headers)
            jobs = self.env["hr.job"].search([("active", "=", True)], limit=10)
            if jobs:
                for j in jobs:
                    ws.append(["Additional Position", j.name, 1, 0, 0, 0, 0, 0, 0, 0, 0, 0, 0, 0, 0, 15000.0, "Operational requirement"])
            else:
                ws.append(["New Position", "Customer Service Officer", 0, 1, 0, 0, 0, 0, 0, 0, 0, 0, 0, 0, 0, 15000.0, "New role"])
        elif cat == "fixed_asset":
            headers = ["Asset Category", "Item Description", "Quantity Requested", "Unit Cost",
                       "Q1 Qty", "Q2 Qty", "Q3 Qty", "Q4 Qty", "Purpose / Justification"]
            ws.append(headers)
            fa_cats = self.env["pbms.fixed.asset.category"].search([("active", "=", True)])
            if fa_cats:
                for fc in fa_cats:
                    ws.append([fc.name or fc.code, f"Sample {fc.name}", 1, 5000.0, 1, 0, 0, 0, "Annual requirement"])
            else:
                ws.append(["Office Equipment", "Desktop Computer with UPS", 2, 85000.0, 2, 0, 0, 0, "New staff workstation"])

        # Format header cells
        for col_idx, cell in enumerate(ws[1], start=1):
            cell.fill = fill_header
            cell.font = font_header
            cell.alignment = align_center

        # Auto column width
        for col in ws.columns:
            max_len = max(len(str(cell.value or "")) for cell in col)
            col_letter = col[0].column_letter
            ws.column_dimensions[col_letter].width = max(max_len + 4, 12)

        buf = io.BytesIO()
        wb.save(buf)
        buf.seek(0)
        file_bytes = buf.getvalue()

        cat_names = dict(self._fields["format_type"].selection)
        file_name = f"PBMS_{cat_names.get(cat, cat).replace(' ', '_')}_Template.xlsx"

        attachment = self.env["ir.attachment"].create({
            "name": file_name,
            "type": "binary",
            "datas": base64.b64encode(file_bytes),
            "mimetype": "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        })
        return {
            "type": "ir.actions.act_url",
            "url": f"/web/content/{attachment.id}?download=true",
            "target": "self",
        }

    def action_import(self):
        """Parse Excel file and import lines into the target plan."""
        self.ensure_one()
        if load_workbook is None:
            raise UserError(_("The 'openpyxl' Python library is required. Ask your administrator to install it."))
        if not self.file_data:
            raise UserError(_("Please attach an .xlsx file to import."))

        wb = load_workbook(io.BytesIO(base64.b64decode(self.file_data)), data_only=True)
        ws = wb.active

        plan = self._get_target_plan()
        if not plan.cycle_id or plan.cycle_id.state != "open":
            cycle_label = dict(plan.cycle_id._fields["state"].selection).get(plan.cycle_id.state, plan.cycle_id.state) if plan.cycle_id else _("Unknown")
            raise UserError(_(
                "Planning cycle '%s' is currently '%s' and is not open for unit input. "
                "You cannot import plan data until the cycle is officially opened by SPPMD."
            ) % (plan.cycle_id.name if plan.cycle_id else "", cycle_label))

        if self.import_mode == "replace":
            old_lines = plan.line_ids.filtered(lambda l: l.line_type == self.format_type)
            if old_lines:
                old_lines.with_context(bypass_plan_lock=True).unlink()

        created, errors = 0, []
        for row_idx, row in enumerate(ws.iter_rows(min_row=2, values_only=True), start=2):
            if not row or not any(row):
                continue
            try:
                if self.format_type == "deposit":
                    created += self._import_deposit_row(row, plan)
                elif self.format_type == "customer_base":
                    created += self._import_customer_base_row(row, plan)
                elif self.format_type == "fx":
                    created += self._import_fx_row(row, plan)
                elif self.format_type == "digital_banking":
                    created += self._import_digital_banking_row(row, plan)
                elif self.format_type == "general_expense":
                    created += self._import_expense_row(row, plan)
                elif self.format_type == "manpower":
                    created += self._import_manpower_row(row, plan)
                elif self.format_type == "fixed_asset":
                    created += self._import_fixed_asset_row(row, plan)
            except Exception as exc:
                errors.append(_("Row %(row)s: %(error)s", row=row_idx, error=str(exc)))

        if errors:
            raise UserError(_(
                "%(created)s row(s) imported, %(failed)s failed:\n%(errors)s",
                created=created, failed=len(errors), errors="\n".join(errors)
            ))

        # Recompute totals and summaries on plan
        if hasattr(plan, "_compute_category_summaries"):
            plan._compute_category_summaries()
        if hasattr(plan, "_compute_plan_notification_details"):
            plan._compute_plan_notification_details()

        return {
            "type": "ir.actions.client",
            "tag": "display_notification",
            "params": {
                "title": _("Import Successful"),
                "message": _("%s row(s) successfully imported into %s.") % (created, plan.display_name),
                "type": "success",
                "next": {"type": "ir.actions.act_window_close"},
            },
        }

    def _get_float(self, val, field_name="Amount"):
        if val is None or val == "" or str(val).strip() == "":
            return 0.0
        val_str = str(val).replace(",", "").strip()
        try:
            return float(val_str)
        except (ValueError, TypeError):
            raise UserError(_("Invalid numeric input '%s' for '%s'. Numbers cannot contain text or character values.") % (val, field_name))

    def _get_int(self, val, field_name="Quantity"):
        if val is None or val == "" or str(val).strip() == "":
            return 0
        val_str = str(val).replace(",", "").strip()
        try:
            f = float(val_str)
            if f != int(f):
                raise UserError(_("Expected a whole number (integer) for '%s', but got '%s'.") % (field_name, val))
            return int(f)
        except (ValueError, TypeError):
            raise UserError(_("Invalid numeric input '%s' for '%s'. Whole numbers cannot contain text or character values.") % (val, field_name))

    def _get_text(self, val, field_name="Text", required=False):
        if val is None or str(val).strip() == "":
            if required:
                raise UserError(_("Field '%s' is required and cannot be empty.") % field_name)
            return ""
        s = str(val).strip()
        cleaned = s.replace(".", "").replace(",", "").replace("-", "").replace("+", "").replace(" ", "")
        if cleaned and cleaned.isdigit():
            raise UserError(_("Field '%s' cannot be purely numbers ('%s'). Please provide a valid text description.") % (field_name, s))
        return s

    def _import_deposit_row(self, row, plan):
        code = str(row[0] or "").strip()
        name = str(row[1] or "").strip() if len(row) > 1 else ""
        deposit_type = self.env["pbms.deposit.type"].search([
            "|", ("code", "=ilike", code), ("name", "=ilike", name or code)
        ], limit=1)
        if not deposit_type:
            raise UserError(_("Unknown deposit product code or name: '%s'") % (code or name))

        ob = self._get_float(row[2], _("Opening Balance")) if len(row) > 2 else 0.0
        months = [self._get_float(row[i], f"Month {i-2}") if len(row) > i else 0.0 for i in range(3, 15)]

        vals = {
            "plan_id": plan.id,
            "line_type": "deposit",
            "source_unit_id": plan.org_unit_id.id,
            "deposit_type_id": deposit_type.id,
            "opening_balance": ob,
        }
        vals.update({f"m{str(i + 1).zfill(2)}": months[i] for i in range(12)})
        self.env["pbms.plan.category.line"].with_context(bypass_plan_lock=True).create(vals)
        return 1

    def _import_customer_base_row(self, row, plan):
        code = str(row[0] or "").strip()
        name = str(row[1] or "").strip() if len(row) > 1 else ""
        deposit_type = self.env["pbms.deposit.type"].search([
            "|", ("code", "=ilike", code), ("name", "=ilike", name or code)
        ], limit=1)
        if not deposit_type:
            raise UserError(_("Unknown customer segment code or name: '%s'") % (code or name))

        ob = self._get_float(row[2], _("Opening Balance")) if len(row) > 2 else 0.0
        months = [self._get_float(row[i], f"Month {i-2}") if len(row) > i else 0.0 for i in range(3, 15)]

        vals = {
            "plan_id": plan.id,
            "line_type": "customer_base",
            "source_unit_id": plan.org_unit_id.id,
            "deposit_type_id": deposit_type.id,
            "opening_balance": ob,
        }
        vals.update({f"m{str(i + 1).zfill(2)}": months[i] for i in range(12)})
        self.env["pbms.plan.category.line"].with_context(bypass_plan_lock=True).create(vals)
        return 1

    def _import_fx_row(self, row, plan):
        code = str(row[0] or "").strip()
        name = str(row[1] or "").strip() if len(row) > 1 else ""
        fx_type = self.env["pbms.fx.source.type"].search([
            "|", ("code", "=ilike", code), ("name", "=ilike", name or code)
        ], limit=1)
        if not fx_type:
            raise UserError(_("Unknown FX source code or name: '%s'") % (code or name))

        months = [self._get_float(row[i], f"Month {i-1}") if len(row) > i else 0.0 for i in range(2, 14)]

        vals = {
            "plan_id": plan.id,
            "line_type": "fx",
            "source_unit_id": plan.org_unit_id.id,
            "fx_source_type": fx_type.id,
        }
        vals.update({f"m{str(i + 1).zfill(2)}": months[i] for i in range(12)})
        self.env["pbms.plan.category.line"].with_context(bypass_plan_lock=True).create(vals)
        return 1

    def _import_digital_banking_row(self, row, plan):
        code = str(row[0] or "").strip()
        name = str(row[1] or "").strip() if len(row) > 1 else ""
        ch = self.env["pbms.digital.channel"].search([
            "|", ("code", "=ilike", code), ("name", "=ilike", name or code)
        ], limit=1)
        if not ch:
            raise UserError(_("Unknown digital channel code or name: '%s'") % (code or name))

        months = [self._get_float(row[i], f"Month {i-1}") if len(row) > i else 0.0 for i in range(2, 14)]

        vals = {
            "plan_id": plan.id,
            "line_type": "digital_banking",
            "source_unit_id": plan.org_unit_id.id,
            "channel_id": ch.id,
        }
        vals.update({f"m{str(i + 1).zfill(2)}": months[i] for i in range(12)})
        self.env["pbms.plan.category.line"].with_context(bypass_plan_lock=True).create(vals)
        return 1

    def _import_expense_row(self, row, plan):
        code = str(row[0] or "").strip()
        name = str(row[1] or "").strip() if len(row) > 1 else ""
        acc = self.env["pbms.expense.account"].search([
            "|", ("code", "=ilike", code), ("name", "=ilike", name or code)
        ], limit=1)
        if not acc:
            raise UserError(_("Unknown expense account code or name: '%s'") % (code or name))

        months = [self._get_float(row[i], f"Month {i-1}") if len(row) > i else 0.0 for i in range(2, 14)]

        vals = {
            "plan_id": plan.id,
            "line_type": "general_expense",
            "source_unit_id": plan.org_unit_id.id,
            "expense_account_id": acc.id,
        }
        vals.update({f"m{str(i + 1).zfill(2)}": months[i] for i in range(12)})
        self.env["pbms.plan.category.line"].with_context(bypass_plan_lock=True).create(vals)
        return 1

    def _import_manpower_row(self, row, plan):
        pos_type_name = self._get_text(row[0], _("Position Type"), required=True)
        job_title = self._get_text(row[1], _("Job Title"), required=True)
        baseline = self._get_int(row[2], _("Baseline")) if len(row) > 2 else 0

        pos_type = self.env["pbms.position.type"].search([
            "|", ("code", "=ilike", pos_type_name), ("name", "=ilike", pos_type_name)
        ], limit=1)
        job = self.env["hr.job"].search([("name", "=ilike", job_title)], limit=1)

        p_code = (
            "new"
            if ((pos_type and (pos_type.is_new_position or pos_type.code in ("new", "new_position")))
                or ("new" in pos_type_name.lower()))
            else "additional"
        )

        if p_code != "new" and not job:
            job = self.env["hr.job"].search([("name", "ilike", job_title)], limit=1)
            if not job:
                raise UserError(_("Job Position '%s' was not found in HR Job Positions. Please make sure the job title exists or select 'New Position'.") % job_title)

        months = [self._get_int(row[i], f"Month {i-2}") if len(row) > i else 0 for i in range(3, 15)]
        monthly_comp = self._get_float(row[15], _("Monthly Compensation")) if len(row) > 15 else 0.0
        reason = self._get_text(row[16], _("Reason")) if len(row) > 16 else ""

        vals = {
            "plan_id": plan.id,
            "line_type": "manpower",
            "source_unit_id": plan.org_unit_id.id,
            "position_type_id": pos_type.id if pos_type else False,
            "position_type": p_code,
            "job_id": job.id if job else False,
            "new_job_title": job_title if (p_code == "new" or not job) else False,
            "base_salary": monthly_comp,
            "other_justification": reason,
        }
        vals.update({f"m{str(i + 1).zfill(2)}": months[i] for i in range(12)})
        self.env["pbms.plan.category.line"].with_context(bypass_plan_lock=True).create(vals)
        return 1

    def _import_fixed_asset_row(self, row, plan):
        cat_name = self._get_text(row[0], _("Asset Category"), required=True)
        desc = self._get_text(row[1], _("Item Description")) if len(row) > 1 else ""
        qty = self._get_int(row[2], _("Quantity")) if len(row) > 2 else 0
        unit_cost = self._get_float(row[3], _("Estimated Unit Price")) if len(row) > 3 else 0.0
        q1 = self._get_int(row[4], _("Q1 Quantity")) if len(row) > 4 else 0
        q2 = self._get_int(row[5], _("Q2 Quantity")) if len(row) > 5 else 0
        q3 = self._get_int(row[6], _("Q3 Quantity")) if len(row) > 6 else 0
        q4 = self._get_int(row[7], _("Q4 Quantity")) if len(row) > 7 else 0
        purpose_str = self._get_text(row[8], _("Purpose")) if len(row) > 8 else ""
        job = False
        if purpose_str:
            job = self.env["hr.job"].search([("name", "=ilike", purpose_str)], limit=1)

        if (q1 + q2 + q3 + q4) == 0 and qty > 0:
            q1 = qty

        fa_cat = self.env["pbms.fixed.asset.category"].search([
            "|", ("code", "=ilike", cat_name), ("name", "=ilike", cat_name)
        ], limit=1)
        if not fa_cat:
            raise UserError(_("Fixed Asset Category '%s' was not found.") % cat_name)

        vals = {
            "plan_id": plan.id,
            "line_type": "fixed_asset",
            "source_unit_id": plan.org_unit_id.id,
            "category_id": fa_cat.id,
            "item_description": desc or cat_name,
            "estimated_unit_price": unit_cost,
            "fa_q1": q1,
            "fa_q2": q2,
            "fa_q3": q3,
            "fa_q4": q4,
            "purpose": job.id if job else False,
        }
        self.env["pbms.plan.category.line"].with_context(bypass_plan_lock=True).create(vals)
        return 1
