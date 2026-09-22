import os
import re
import sys
import openpyxl
import psycopg2
import odoo
from odoo.modules.registry import Registry
from odoo import api, SUPERUSER_ID

BASE_DIR = "/mnt/extra-addons/competency_management/data/comptency_files"
FILE_COMP = os.path.join(BASE_DIR, "hr.competency (1).xlsx")
FILE_REQ = os.path.join(BASE_DIR, "hr.competency.requirement.xlsx")

# Hand-crafted definitions and indicators for Core and Leadership
CORE_INDICATORS = {
    'collaboration': {
        'def': "Actively works across boundaries, builds inclusive partnerships, shares knowledge, and leverages diverse perspectives to achieve Bunna Bank's collective goals.",
        '1': "Cooperates effectively with immediate team members; shares information openly; listens attentively and respects diverse viewpoints.",
        '2': "Proactively coordinates across departmental teams; resolves interpersonal friction constructively; offers assistance during peak operational workloads.",
        '3': "Facilitates cross-functional collaboration and joint problem-solving; breaks down organizational silos; establishes collaborative operational routines.",
        '4': "Fosters an enterprise-wide culture of partnership and mutual accountability; establishes strategic alliances and collaborative frameworks across all bank directorates."
    },
    'creativity': {
        'def': "Generates novel, value-adding ideas, embraces innovation, and finds resourceful solutions to enhance banking operations and customer experiences.",
        '1': "Suggests practical improvements to routine personal tasks; open to adopting new tools and digital processes.",
        '2': "Identifies workflow bottlenecks and proposes innovative, workable solutions; experiments with new ways of delivering customer value.",
        '3': "Designs and leads departmental process innovations; applies creative problem-solving to recurring operational hurdles; champions change initiatives.",
        '4': "Drives bank-wide innovation strategy and digital transformation; cultivates an organizational environment that rewards agile experimentation and creative thinking."
    },
    'ethical_influence': {
        'def': "Influences others through high ethical integrity, persuasive communication, transparent decision-making, and unwavering adherence to banking compliance.",
        '1': "Demonstrates personal honesty, transparency, and respect in all daily interactions; adheres strictly to bank code of conduct.",
        '2': "Persuades colleagues and customers using facts and ethical reasoning; addresses and reports compliance concerns constructively; maintains confidential information.",
        '3': "Builds widespread consensus for ethical practices; navigates sensitive stakeholder discussions with diplomatic integrity; mentors peers on compliance dilemmas.",
        '4': "Shapes enterprise-level ethical standards and corporate governance; serves as a trusted advisor to senior executive leadership; champions institutional trust."
    },
    'execution_mastery': {
        'def': "Consistently delivers outstanding, timely, and error-free operational results through disciplined planning, relentless focus on quality, and accountability.",
        '1': "Prioritizes daily assignments effectively; meets deadlines reliably; maintains accuracy in routine transaction processing and documentation.",
        '2': "Manages multiple deliverables independently under pressure; anticipates obstacles and takes corrective actions; consistently meets or exceeds performance targets.",
        '3': "Drives end-to-end departmental operational excellence; establishes stringent quality benchmarks; ensures flawless execution of complex bank initiatives.",
        '4': "Establishes bank-wide execution excellence frameworks; ensures organizational agility, accountability, and disciplined delivery of multi-year strategic objectives."
    },
    'professional_authenticity': {
        'def': "Exemplifies genuine self-awareness, personal accountability, professionalism, and commitment to continuous growth as a representative of Bunna Bank.",
        '1': "Maintains a professional, courteous demeanor; accepts constructive feedback gracefully; takes responsibility for personal actions and mistakes.",
        '2': "Demonstrates consistent alignment between personal values and bank standards; handles stressful situations calmly and professionally; actively seeks self-development.",
        '3': "Inspires confidence and credibility among internal and external stakeholders; demonstrates high emotional intelligence; coaches others on professional presence.",
        '4': "Embodies and personifies Bunna Bank's brand, culture, and values at an institutional level; commands broad respect across the banking sector."
    }
}

LEADERSHIP_INDICATORS = {
    'ambidextrous_leadership': {
        'def': "Balances operational excellence and daily efficiency with forward-looking exploration, strategic adaptability, and business growth.",
        '1': "Executes current operational procedures efficiently while showing openness to organizational change and new operational practices.",
        '2': "Balances daily service delivery targets with the implementation of departmental improvements and transformation milestones.",
        '3': "Effectively allocates resources between steady-state banking operations and innovation projects; drives operational efficiency while capturing new business opportunities.",
        '4': "Leads enterprise-wide strategic agility; successfully steers the bank through digital disruption while maintaining robust operational stability and profitability."
    },
    'continuous_improvement': {
        'def': "Relentlessly pursues operational enhancement, cost optimization, service turnaround acceleration, and best-in-class banking practices.",
        '1': "Participates actively in team reviews; identifies recurring errors or delays in daily tasks and suggests remedies.",
        '2': "Analyzes unit performance data to streamline workflows; eliminates non-value-adding steps; implements customer turnaround time reductions.",
        '3': "Champions continuous improvement projects across the division; measures and reports measurable productivity and quality gains.",
        '4': "Institutionalizes an enterprise-wide culture of continuous excellence; benchmarks Bunna Bank against global banking standards; drives sustained competitive advantages."
    },
    'people_leadership': {
        'def': "Inspires, develops, empowers, and aligns teams to achieve peak performance, fostering inclusion, engagement, and talent retention.",
        '1': "Supports team members in onboarding and daily routines; provides clear peer guidance and fosters an inclusive environment.",
        '2': "Sets clear performance expectations and targets; provides timely, constructive coaching and feedback; recognizes and rewards strong performance.",
        '3': "Builds high-performing teams; identifies and develops future leadership talent; manages employee engagement and resolves complex team dynamics.",
        '4': "Shapes enterprise human capital strategy; champions leadership succession pipelines; cultivates an employer-of-choice reputation for Bunna Bank."
    },
    'prudential_decision_making': {
        'def': "Makes sound, timely, and risk-conscious commercial decisions grounded in rigorous data analysis, regulatory compliance, and bank stewardship.",
        '1': "Applies bank guidelines and compliance rules accurately to routine decisions; escalates ambiguous or high-risk situations appropriately.",
        '2': "Evaluates risks, costs, and benefits systematically before deciding; exercises sound judgment in standard credit, operational, and customer matters.",
        '3': "Makes difficult, high-impact operational and risk decisions in complex, uncertain environments; balances risk mitigation with commercial viability.",
        '4': "Makes enterprise-defining strategic decisions affecting bank solvency, asset-liability management, and long-term shareholder value; maintains strong regulatory standing."
    },
    'self_leadership': {
        'def': "Demonstrates proactive initiative, emotional resilience, personal accountability, and continuous commitment to self-mastery.",
        '1': "Takes ownership of daily tasks without requiring constant supervision; manages time effectively; demonstrates positive attitude under pressure.",
        '2': "Demonstrates resilience during organizational change; sets personal development goals; takes initiative to acquire new banking competencies.",
        '3': "Serves as an exemplar of self-discipline and emotional maturity; effectively manages personal stress in high-stakes situations; pursues continuous executive learning.",
        '4': "Demonstrates exceptional visionary self-mastery; leads by personal example under intense scrutiny; projects inspiring calm, clarity, and determination."
    },
    'strategy_management': {
        'def': "Formulates, aligns, communicates, and executes strategic roadmaps that drive sustainable market share growth, profitability, and customer value.",
        '1': "Understands the bank's strategic vision and how daily unit outputs contribute to bank-wide objectives.",
        '2': "Translates departmental strategy into actionable quarterly milestones and KPI scorecards; monitors progress and adjusts tactics as needed.",
        '3': "Conducts environmental scans and competitive analyses; devises robust divisional strategic plans; aligns resources with high-return business opportunities.",
        '4': "Formulates enterprise strategy for Bunna Bank; drives corporate growth, market expansion, and digital ecosystem leadership; creates enduring shareholder value."
    }
}

def generate_technical_indicators(name, jf):
    c_name = name.strip()
    jf_str = f" in {jf}" if jf else ""
    definition = f"Demonstrates specialized technical expertise, analytical proficiency, and operational mastery in {c_name}{jf_str} to safeguard bank assets and optimize operational outcomes."
    
    indicators = {
        '1': f"Understands fundamental principles, core workflows, and standard operating procedures of {c_name}; performs foundational tasks under routine supervision.",
        '2': f"Independently executes key operational tasks, methodologies, and technical tools relating to {c_name}; resolves routine operational exceptions with accuracy.",
        '3': f"Applies advanced specialized knowledge in {c_name} to address complex cases; optimizes workflow efficiency, troubleshoots critical incidents, and coaches colleagues.",
        '4': f"Serves as institutional authority and subject-matter expert in {c_name}; establishes departmental governance standards, drives innovation, and ensures strict regulatory compliance."
    }
    return definition, indicators


def run_seeding(db_name="ERP_TEST3", dry_run=False):
    print(f"[*] Connecting to Odoo registry for database '{db_name}'...")
    reg = Registry(db_name)
    with reg.cursor() as cr:
        env = api.Environment(cr, SUPERUSER_ID, {})

        CompObj = env['competency.competency']
        LevelObj = env['competency.proficiency.level']
        RatingModelObj = env['competency.rating.model']
        JobObj = env['hr.job']
        MapObj = env['competency.role.mapping']
        MapLineObj = env['competency.role.mapping.line']

        scale_4 = RatingModelObj.search([('code', '=', '4SCALE')], limit=1)
        if not scale_4:
            scale_4 = RatingModelObj.search([], limit=1)

        # ----------------------------------------------------------------------
        # STEP 1: SEED COMPETENCY DICTIONARY (541 Competencies)
        # ----------------------------------------------------------------------
        print("\n" + "=" * 60)
        print("[*] STEP 1: SEEDING COMPETENCY DICTIONARY")
        print("=" * 60)
        wb_comp = openpyxl.load_workbook(FILE_COMP, data_only=True)
        ws_comp = wb_comp.active

        existing_comps = CompObj.with_context(active_test=False).search([])
        comp_by_code = {c.code.strip().lower(): c for c in existing_comps if c.code}
        comp_by_name = {c.name.strip().lower(): c for c in existing_comps if c.name}

        created_comps = 0
        updated_comps = 0

        level_names = {
            '1': 'Level 1 - Basic',
            '2': 'Level 2 - Intermediate',
            '3': 'Level 3 - Advanced',
            '4': 'Level 4 - Expert'
        }

        for r_idx in range(2, ws_comp.max_row + 1):
            c_name = ws_comp.cell(row=r_idx, column=1).value
            c_code = ws_comp.cell(row=r_idx, column=2).value
            pillar_raw = ws_comp.cell(row=r_idx, column=3).value
            jf = ws_comp.cell(row=r_idx, column=4).value

            if not c_name or not c_code:
                continue

            name = str(c_name).strip()
            code = str(c_code).strip()
            pillar_str = str(pillar_raw or '').lower()
            if 'lead' in pillar_str:
                pillar = 'leadership'
            elif 'tech' in pillar_str:
                pillar = 'technical'
            else:
                pillar = 'core'

            clean_jf = str(jf).strip() if jf else ('Corporate Core' if pillar == 'core' else ('Bank Leadership' if pillar == 'leadership' else 'General Banking'))

            # Get definition and indicators
            code_key = code.lower().replace('-', '_')
            if pillar == 'core' and code_key in CORE_INDICATORS:
                definition = CORE_INDICATORS[code_key]['def']
                lvl_indicators = {k: CORE_INDICATORS[code_key][k] for k in ['1', '2', '3', '4']}
            elif pillar == 'leadership' and code_key in LEADERSHIP_INDICATORS:
                definition = LEADERSHIP_INDICATORS[code_key]['def']
                lvl_indicators = {k: LEADERSHIP_INDICATORS[code_key][k] for k in ['1', '2', '3', '4']}
            else:
                definition, lvl_indicators = generate_technical_indicators(name, clean_jf)

            comp = comp_by_code.get(code.lower()) or comp_by_name.get(name.lower())

            comp_vals = {
                'name': name,
                'code': code,
                'pillar': pillar,
                'functional_domain': clean_jf,
                'definition': definition,
                'rating_model_id': scale_4.id if scale_4 else False,
                'state': 'approved',
                'active': True,
            }

            if not comp:
                # Build levels
                levels = []
                for lvl_k in ['1', '2', '3', '4']:
                    levels.append((0, 0, {
                        'level': lvl_k,
                        'name': level_names[lvl_k],
                        'definition': f"{level_names[lvl_k]} for {name}",
                        'behavioral_indicators': lvl_indicators[lvl_k],
                    }))
                comp_vals['proficiency_level_ids'] = levels
                comp = CompObj.create(comp_vals)
                comp_by_code[code.lower()] = comp
                comp_by_name[name.lower()] = comp
                created_comps += 1
            else:
                comp.write(comp_vals)
                # Ensure 4 levels exist
                for lvl_k in ['1', '2', '3', '4']:
                    lvl_rec = LevelObj.search([('competency_id', '=', comp.id), ('level', '=', lvl_k)], limit=1)
                    if lvl_rec:
                        lvl_rec.write({
                            'name': level_names[lvl_k],
                            'behavioral_indicators': lvl_indicators[lvl_k],
                        })
                    else:
                        LevelObj.create({
                            'competency_id': comp.id,
                            'level': lvl_k,
                            'name': level_names[lvl_k],
                            'definition': f"{level_names[lvl_k]} for {name}",
                            'behavioral_indicators': lvl_indicators[lvl_k],
                        })
                comp_by_code[code.lower()] = comp
                comp_by_name[name.lower()] = comp
                updated_comps += 1

        print(f"[✔] Competencies Processed: {created_comps + updated_comps} (Created: {created_comps}, Updated: {updated_comps})")

        # ----------------------------------------------------------------------
        # STEP 2: SEED ROLE MAPPINGS (7,760 Requirements -> 522 Role Profiles)
        # ----------------------------------------------------------------------
        print("\n" + "=" * 60)
        print("[*] STEP 2: SEEDING ROLE-COMPETENCY MAPPINGS")
        print("=" * 60)
        wb_req = openpyxl.load_workbook(FILE_REQ, data_only=True)
        ws_req = wb_req.active

        # Cache DB jobs
        all_jobs = JobObj.search([])
        job_by_name = {j.name.strip().lower(): j for j in all_jobs if j.name}

        # Handle specific aliases
        job_aliases = {
            'acting/manager-work force planning and organizational design division': 'manager - work force planning and organizational design division',
            'senior digital marketing officer': 'digital marketing officer'
        }

        # Group requirements by Job Position
        job_requirements = {}
        for r_idx in range(2, ws_req.max_row + 1):
            j_name = ws_req.cell(row=r_idx, column=1).value
            c_name = ws_req.cell(row=r_idx, column=3).value
            lvl_raw = ws_req.cell(row=r_idx, column=5).value

            if not j_name or not c_name:
                continue

            j_str = str(j_name).strip()
            c_str = str(c_name).strip()
            lvl_str = str(lvl_raw or '').strip()
            req_lvl = '2'
            if '1' in lvl_str: req_lvl = '1'
            elif '2' in lvl_str: req_lvl = '2'
            elif '3' in lvl_str: req_lvl = '3'
            elif '4' in lvl_str: req_lvl = '4'

            if j_str not in job_requirements:
                job_requirements[j_str] = {}
            # Deduplicate per job position
            job_requirements[j_str][c_str] = req_lvl

        print(f"[*] Total unique Job Positions to map: {len(job_requirements)}")

        created_maps = 0
        updated_maps = 0
        total_mapped_lines = 0

        for j_str, reqs in job_requirements.items():
            lookup_name = j_str.lower().strip()
            if lookup_name in job_aliases:
                lookup_name = job_aliases[lookup_name]

            job = job_by_name.get(lookup_name)
            if not job:
                # If still not found, create the job in hr.job
                print(f"  [+] Creating missing job position in hr.job: '{j_str}'")
                job = JobObj.create({'name': j_str, 'active': True})
                job_by_name[lookup_name] = job

            # Find or create role mapping for this job
            mapping = MapObj.search([('job_position_id', '=', job.id), ('state', 'in', ['approved', 'under_approval', 'draft'])], limit=1)
            
            if not mapping:
                # Build all lines
                line_vals = []
                for comp_name, req_lvl in reqs.items():
                    comp = comp_by_name.get(comp_name.lower()) or comp_by_code.get(comp_name.lower())
                    if not comp:
                        continue
                    line_vals.append((0, 0, {
                        'competency_id': comp.id,
                        'required_proficiency': req_lvl,
                        'weight': 1.0,
                        'override_default': True,
                    }))
                
                mapping = MapObj.create({
                    'job_position_id': job.id,
                    'version': 'v1.0',
                    'state': 'approved',
                    'line_ids': line_vals
                })
                created_maps += 1
                total_mapped_lines += len(line_vals)
            else:
                # Update existing mapping
                existing_lines = {l.competency_id.id: l for l in mapping.line_ids if l.competency_id}
                new_line_vals = []
                for comp_name, req_lvl in reqs.items():
                    comp = comp_by_name.get(comp_name.lower()) or comp_by_code.get(comp_name.lower())
                    if not comp:
                        continue
                    if comp.id in existing_lines:
                        existing_lines[comp.id].write({
                            'required_proficiency': req_lvl,
                            'override_default': True,
                        })
                    else:
                        new_line_vals.append((0, 0, {
                            'competency_id': comp.id,
                            'required_proficiency': req_lvl,
                            'weight': 1.0,
                            'override_default': True,
                        }))
                if new_line_vals:
                    mapping.write({'line_ids': new_line_vals})
                updated_maps += 1
                total_mapped_lines += len(reqs)

        print(f"\n[✔] Role Mappings Processed: {created_maps + updated_maps}")
        print(f"    - New Mappings Created: {created_maps}")
        print(f"    - Existing Mappings Updated: {updated_maps}")
        print(f"    - Total Competency Lines Linked: {total_mapped_lines}")

        if dry_run:
            print("\n[!] DRY RUN: Rolling back transaction (No changes committed to database).")
            cr.rollback()
        else:
            cr.commit()
            print("\n[✔] SUCCESS: All changes successfully committed to database 'ERP_TEST3'!")

if __name__ == "__main__":
    dry_run_flag = "--dry-run" in sys.argv
    run_seeding("ERP_TEST3", dry_run=dry_run_flag)
