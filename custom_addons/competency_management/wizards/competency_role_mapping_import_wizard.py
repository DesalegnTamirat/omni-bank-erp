# -*- coding: utf-8 -*-
import base64
import io
import re
import openpyxl
import xlsxwriter
from odoo import api, fields, models, _
from odoo.exceptions import UserError, ValidationError


class CompetencyRoleMappingImportWizard(models.TransientModel):
    """Excel Import Wizard for Role-Competency Mappings with template generator,
    pre-validation, and support for versioning and status selection.
    """
    _name = 'competency.role.mapping.import.wizard'
    _description = 'Import Role Mappings from Excel'

    file = fields.Binary(string='Excel File (.xlsx)')
    filename = fields.Char(string='Filename', default='role_mappings.xlsx')
    import_mode = fields.Selection([
        ('new_version', 'Create New Version & Archive Prior (Recommended)'),
        ('override', 'Override Existing Mapping in Place'),
    ], string='Import Mode', default='new_version', required=True,
       help="New Version: Archives existing approved mapping and creates a new incremented version (e.g. v1.1).\n"
            "Override: Replaces competency lines directly on the active mapping.")

    import_state = fields.Selection([
        ('draft', 'All as Draft (Pending Review)'),
        ('approved', 'All as Approved'),
        ('from_excel', 'Use Status from Excel (Admin Only, defaults to Draft)'),
    ], string='Mapping Status', default='draft', required=True,
       help="Draft: Mappings require review and manual approval.\n"
            "Approved: Mappings are immediately finalized and activated.\n"
            "Note: Only Competency Administrators can set mappings directly to Approved.")

    is_admin = fields.Boolean(string='Is Admin', compute='_compute_is_admin')
    state = fields.Selection([
        ('draft', 'Upload & Configure'),
        ('done', 'Completed'),
    ], string='State', default='draft')

    result_summary = fields.Html(string='Summary & Log', readonly=True)
    created_mapping_ids = fields.Many2many(
        'competency.role.mapping',
        'comp_role_mapping_import_wizard_rel',
        'wizard_id', 'mapping_id',
        string='Processed Mappings', readonly=True)

    @api.depends_context('uid')
    def _compute_is_admin(self):
        is_admin = self.env.user.has_group('competency_management.group_competency_admin') or self.env.su
        for rec in self:
            rec.is_admin = is_admin

    def action_download_template(self):
        """Generates a styled .xlsx template with data validation and reference lists."""
        output = io.BytesIO()
        workbook = xlsxwriter.Workbook(output, {'in_memory': True})

        # Styling Formats
        fmt_title = workbook.add_format({
            'bold': True, 'font_size': 14, 'font_color': '#FFFFFF',
            'bg_color': '#541718', 'align': 'center', 'valign': 'vcenter'
        })
        fmt_section = workbook.add_format({
            'bold': True, 'font_size': 11, 'font_color': '#726732',
            'bg_color': '#F8F9FA', 'bottom': 2, 'bottom_color': '#726732'
        })
        fmt_header = workbook.add_format({
            'bold': True, 'font_size': 10, 'font_color': '#FFFFFF',
            'bg_color': '#541718', 'align': 'center', 'valign': 'vcenter',
            'border': 1, 'border_color': '#dee2e6'
        })
        fmt_sub_header = workbook.add_format({
            'bold': True, 'font_size': 9, 'font_color': '#FFFFFF',
            'bg_color': '#726732', 'align': 'center', 'valign': 'vcenter',
            'border': 1
        })
        fmt_cell = workbook.add_format({
            'font_size': 10, 'border': 1, 'border_color': '#E2E8F0', 'valign': 'vcenter'
        })
        fmt_cell_center = workbook.add_format({
            'font_size': 10, 'border': 1, 'border_color': '#E2E8F0', 'align': 'center', 'valign': 'vcenter'
        })
        fmt_bold_label = workbook.add_format({'bold': True, 'font_size': 10})
        fmt_text_desc = workbook.add_format({'font_size': 10, 'text_wrap': True})

        # -------------------------------------------------------------
        # SHEET 1: Instructions
        # -------------------------------------------------------------
        ws_inst = workbook.add_worksheet('Instructions')
        ws_inst.set_column('A:A', 5)
        ws_inst.set_column('B:B', 24)
        ws_inst.set_column('C:C', 75)

        ws_inst.merge_range('B2:C2', 'BUNNA BANK - ROLE COMPETENCY MAPPING IMPORT GUIDE', fmt_title)
        ws_inst.set_row(1, 35)

        ws_inst.write('B4', 'Step 1', fmt_bold_label)
        ws_inst.write('C4', 'Switch to the "Role Mapping Template" tab to enter your job role mapping requirements.', fmt_text_desc)

        ws_inst.write('B5', 'Step 2', fmt_bold_label)
        ws_inst.write('C5', 'Use the "Reference Data" tab to view valid Job Positions, Competency Names/Codes, and Operating Units.', fmt_text_desc)

        ws_inst.write('B6', 'Step 3', fmt_bold_label)
        ws_inst.write('C6', 'Save your filled Excel file and upload it in the ERP Role Mapping Import Wizard.', fmt_text_desc)

        ws_inst.write('B8', 'COLUMN SPECIFICATIONS', fmt_section)
        ws_inst.write('C8', '', fmt_section)

        specs = [
            ('Job Position *', 'Mandatory. Exact Job Position name or code as defined in Bunna Bank ERP.'),
            ('Operating Unit (Optional)', 'Optional. Operating Unit (Branch / Department / District). Leave BLANK for bank-wide global mapping across all branches.'),
            ('Competency *', 'Mandatory. Exact Competency Name or Competency Code (e.g. CORE-01, Execution Mastery).'),
            ('Required Level (1-4) *', 'Mandatory. Select proficiency required: 1 (Basic), 2 (Intermediate), 3 (Advanced), 4 (Expert).'),
            ('Weight', 'Optional. Numerical weight for gap analysis (e.g. 1.0, 1.5, 2.0). Defaults to 1.0 if empty.'),
            ('Override Default', 'Optional. TRUE or FALSE. Set to TRUE if this level deliberately overrides matrix defaults. Defaults to FALSE.'),
            ('Status (Draft/Approved)', 'Optional. Draft or Approved. Note: Direct approval is only available when imported by Competency Administrators.'),
        ]

        row = 9
        for col_name, desc in specs:
            ws_inst.write(row, 1, col_name, fmt_bold_label)
            ws_inst.write(row, 2, desc, fmt_text_desc)
            ws_inst.set_row(row, 22)
            row += 1

        ws_inst.write(row + 1, 1, 'IMPORTANT RULES', fmt_section)
        ws_inst.write(row + 1, 2, '', fmt_section)
        rules = [
            '1. Each job position can have multiple competencies by repeating the Job Position on multiple rows.',
            '2. Do NOT duplicate the same competency under the same Job Position.',
            '3. In "Create New Version" mode, existing approved mappings will be automatically archived and a new version will be created.',
            '4. In "Override" mode, existing mapping lines are replaced in place without archiving.',
        ]
        r_idx = row + 2
        for r in rules:
            ws_inst.write(r_idx, 1, 'Rule', fmt_bold_label)
            ws_inst.write(r_idx, 2, r, fmt_text_desc)
            r_idx += 1

        # -------------------------------------------------------------
        # SHEET 2: Role Mapping Template
        # -------------------------------------------------------------
        ws_tmpl = workbook.add_worksheet('Role Mapping Template')
        headers = [
            ('Job Position', 32),
            ('Operating Unit (Optional)', 26),
            ('Competency', 36),
            ('Required Level (1-4)', 22),
            ('Weight', 12),
            ('Override Default', 18),
            ('Status (Draft/Approved)', 22),
        ]

        for i, (h, width) in enumerate(headers):
            ws_tmpl.write(0, i, h, fmt_header)
            ws_tmpl.set_column(i, i, width)
        ws_tmpl.set_row(0, 28)

        # Query sample positions & competencies for demonstration
        sample_jobs = self.env['hr.job'].search([], limit=1)
        sample_comps = self.env['competency.competency'].search([('state', '=', 'approved'), ('status', '=', 'active')], limit=3)

        job_demo_name = sample_jobs[0].name if sample_jobs else 'Accountant - I'
        comp_demo_1 = sample_comps[0].name if len(sample_comps) > 0 else 'Execution Mastery'
        comp_demo_2 = sample_comps[1].name if len(sample_comps) > 1 else 'Professional Authenticity'
        comp_demo_3 = sample_comps[2].name if len(sample_comps) > 2 else 'Collaboration'

        sample_rows = [
            [job_demo_name, '', comp_demo_1, 3, 1.0, 'FALSE', 'Draft'],
            [job_demo_name, '', comp_demo_2, 2, 1.0, 'FALSE', 'Draft'],
            [job_demo_name, '', comp_demo_3, 2, 1.0, 'FALSE', 'Draft'],
        ]

        for r_i, r_data in enumerate(sample_rows, start=1):
            for c_i, val in enumerate(r_data):
                if c_i in (3, 4):
                    ws_tmpl.write(r_i, c_i, val, fmt_cell_center)
                elif c_i in (5, 6):
                    ws_tmpl.write(r_i, c_i, val, fmt_cell_center)
                else:
                    ws_tmpl.write(r_i, c_i, val, fmt_cell)
            ws_tmpl.set_row(r_i, 20)

        # Add Data Validation on Template
        ws_tmpl.data_validation('D2:D5000', {
            'validate': 'list',
            'source': ['1', '2', '3', '4'],
            'input_title': 'Proficiency Level',
            'input_message': 'Choose 1 (Basic), 2 (Intermediate), 3 (Advanced), or 4 (Expert)',
            'error_title': 'Invalid Level',
            'error_message': 'Level must be 1, 2, 3, or 4.',
        })
        ws_tmpl.data_validation('F2:F5000', {
            'validate': 'list',
            'source': ['TRUE', 'FALSE'],
            'input_title': 'Override Default',
            'input_message': 'Select TRUE if overriding Matrix Default, otherwise FALSE.',
        })
        ws_tmpl.data_validation('G2:G5000', {
            'validate': 'list',
            'source': ['Draft', 'Approved'],
            'input_title': 'Status',
            'input_message': 'Select Draft or Approved (Approval requires Administrator rights).',
        })

        # -------------------------------------------------------------
        # SHEET 3: Reference Data
        # -------------------------------------------------------------
        ws_ref = workbook.add_worksheet('Reference Data')
        ws_ref.set_column('A:B', 32)
        ws_ref.set_column('C:C', 5)
        ws_ref.set_column('D:F', 30)
        ws_ref.set_column('G:G', 5)
        ws_ref.set_column('H:I', 28)

        # Header titles
        ws_ref.merge_range('A1:B1', 'ACTIVE JOB POSITIONS', fmt_sub_header)
        ws_ref.write('A2', 'Job Position Name', fmt_header)
        ws_ref.write('B2', 'Department', fmt_header)

        ws_ref.merge_range('D1:F1', 'APPROVED COMPETENCIES', fmt_sub_header)
        ws_ref.write('D2', 'Competency Name', fmt_header)
        ws_ref.write('E2', 'Code', fmt_header)
        ws_ref.write('F2', 'Pillar', fmt_header)

        ws_ref.merge_range('H1:I1', 'OPERATING UNITS', fmt_sub_header)
        ws_ref.write('H2', 'Operating Unit Name', fmt_header)
        ws_ref.write('I2', 'Code', fmt_header)

        # Populate Reference Jobs
        all_jobs = self.env['hr.job'].search([], order='name asc', limit=500)
        for idx, j in enumerate(all_jobs, start=2):
            ws_ref.write(idx, 0, j.name or '', fmt_cell)
            ws_ref.write(idx, 1, j.department_id.name if j.department_id else '', fmt_cell)

        # Populate Reference Competencies
        all_comps = self.env['competency.competency'].search(
            [('state', '=', 'approved'), ('status', '=', 'active')],
            order='pillar asc, name asc', limit=500
        )
        for idx, c in enumerate(all_comps, start=2):
            ws_ref.write(idx, 3, c.name or '', fmt_cell)
            ws_ref.write(idx, 4, c.code or '', fmt_cell)
            ws_ref.write(idx, 5, dict(c._fields['pillar'].selection).get(c.pillar, c.pillar or ''), fmt_cell)

        # Populate Reference Operating Units (if model exists)
        if 'operating.unit' in self.env:
            all_ous = self.env['operating.unit'].search([('active', '=', True)], order='name asc', limit=500)
            for idx, ou in enumerate(all_ous, start=2):
                ws_ref.write(idx, 7, ou.name or '', fmt_cell)
                ou_code_str = ''
                if ou.code:
                    ou_code_str = str(getattr(ou.code, 'display_name', False) or getattr(ou.code, 'name', False) or '')
                ws_ref.write(idx, 8, ou_code_str, fmt_cell)

        workbook.close()
        xlsx_data = output.getvalue()
        output.close()

        attachment = self.env['ir.attachment'].create({
            'name': 'Role_Competency_Mapping_Template.xlsx',
            'datas': base64.b64encode(xlsx_data),
            'mimetype': 'application/vnd.openxmlformats-officedocument.spreadsheetml.sheet',
        })

        return {
            'type': 'ir.actions.act_url',
            'url': f'/web/content/{attachment.id}?download=true',
            'target': 'self',
        }

    def action_import_file(self):
        """Validates and processes the uploaded Excel file into Role Mappings."""
        self.ensure_one()
        if not self.file:
            raise UserError(_("Please upload an Excel (.xlsx) file to import."))

        try:
            file_bytes = base64.b64decode(self.file)
            wb = openpyxl.load_workbook(io.BytesIO(file_bytes), data_only=True)
        except Exception as e:
            raise UserError(_("Invalid or corrupted Excel file: %s") % str(e))

        # Identify target worksheet
        sheet = None
        for candidate_name in ['Role Mapping Template', 'Template', 'Role Mapping', 'Role_Mapping', 'Sheet1']:
            if candidate_name in wb.sheetnames:
                sheet = wb[candidate_name]
                break
        if sheet is None:
            sheet = wb.active

        # Find header row
        header_row_idx = None
        col_map = {}
        for r_idx, row in enumerate(sheet.iter_rows(values_only=True), start=1):
            row_str = [str(cell).strip().lower() if cell is not None else '' for cell in row]
            has_job = any('job' in c or 'position' in c or 'role' in c for c in row_str)
            has_comp = any('competency' in c for c in row_str)
            if has_job and has_comp:
                header_row_idx = r_idx
                for c_idx, val in enumerate(row_str):
                    if 'job' in val or 'position' in val or 'role' in val:
                        col_map['job'] = c_idx
                    elif 'operating' in val or 'unit' in val or 'branch' in val:
                        col_map['ou'] = c_idx
                    elif 'competency' in val:
                        col_map['competency'] = c_idx
                    elif 'level' in val or 'proficiency' in val:
                        col_map['level'] = c_idx
                    elif 'weight' in val:
                        col_map['weight'] = c_idx
                    elif 'override' in val:
                        col_map['override'] = c_idx
                    elif 'status' in val or 'state' in val:
                        col_map['status'] = c_idx
                break

        if header_row_idx is None or 'job' not in col_map or 'competency' not in col_map or 'level' not in col_map:
            raise UserError(_(
                "Unable to identify header columns. The Excel sheet must contain columns for "
                "'Job Position', 'Competency', and 'Required Level'."
            ))

        # Caches for rapid resolution
        JobModel = self.env['hr.job'].sudo()
        CompModel = self.env['competency.competency'].sudo()
        OuModel = self.env['operating.unit'].sudo() if 'operating.unit' in self.env else None

        job_cache = {}
        comp_cache = {}
        ou_cache = {}

        errors = []
        valid_rows = []
        seen_job_comp = {}

        total_rows_read = 0
        for row_idx, row in enumerate(sheet.iter_rows(min_row=header_row_idx + 1, values_only=True), start=header_row_idx + 1):
            # Check for completely empty row
            if not row or all(c is None or str(c).strip() == '' for c in row):
                continue
            total_rows_read += 1

            # Extract cell values safely
            raw_job = str(row[col_map['job']]).strip() if col_map.get('job') is not None and col_map['job'] < len(row) and row[col_map['job']] is not None else ''
            raw_ou = str(row[col_map['ou']]).strip() if col_map.get('ou') is not None and col_map['ou'] < len(row) and row[col_map['ou']] is not None else ''
            raw_comp = str(row[col_map['competency']]).strip() if col_map.get('competency') is not None and col_map['competency'] < len(row) and row[col_map['competency']] is not None else ''
            raw_lvl = str(row[col_map['level']]).strip() if col_map.get('level') is not None and col_map['level'] < len(row) and row[col_map['level']] is not None else ''
            raw_weight = row[col_map['weight']] if col_map.get('weight') is not None and col_map['weight'] < len(row) and row[col_map['weight']] is not None else 1.0
            raw_override = row[col_map['override']] if col_map.get('override') is not None and col_map['override'] < len(row) and row[col_map['override']] is not None else False
            raw_status = str(row[col_map['status']]).strip() if col_map.get('status') is not None and col_map['status'] < len(row) and row[col_map['status']] is not None else ''

            if not raw_job:
                errors.append({'row': row_idx, 'field': 'Job Position', 'message': _("Job Position cannot be empty.")})
                continue
            if not raw_comp:
                errors.append({'row': row_idx, 'field': 'Competency', 'message': _("Competency cannot be empty.")})
                continue

            # 1. Resolve Job Position
            job = job_cache.get(raw_job.lower())
            if not job:
                job = JobModel.search([('name', '=ilike', raw_job)], limit=1)
                if not job:
                    job = JobModel.search([('name', 'ilike', raw_job)], limit=1)
                if job:
                    job_cache[raw_job.lower()] = job
                else:
                    errors.append({'row': row_idx, 'field': 'Job Position', 'message': _("Job Position '%s' was not found.") % raw_job})
                    continue

            # 2. Resolve Operating Unit (if specified)
            ou = False
            if raw_ou and raw_ou.lower() not in ('none', 'n/a', 'all', 'global', '-', ''):
                if OuModel is not None:
                    ou = ou_cache.get(raw_ou.lower())
                    if not ou:
                        ou = OuModel.search([('name', '=ilike', raw_ou)], limit=1)
                        if not ou:
                            try:
                                ou = OuModel.search([('code.name', '=ilike', raw_ou)], limit=1)
                            except Exception:
                                ou = False
                        if not ou:
                            ou = OuModel.search([('name', 'ilike', raw_ou)], limit=1)
                        if ou:
                            ou_cache[raw_ou.lower()] = ou
                        else:
                            errors.append({'row': row_idx, 'field': 'Operating Unit', 'message': _("Operating Unit '%s' was not found.") % raw_ou})
                            continue

            # 3. Resolve Competency
            comp = comp_cache.get(raw_comp.lower())
            if not comp:
                comp = CompModel.search(['|', ('code', '=ilike', raw_comp), ('name', '=ilike', raw_comp)], limit=1)
                if comp:
                    if comp.state != 'approved' or comp.status != 'active':
                        errors.append({'row': row_idx, 'field': 'Competency', 'message': _("Competency '%s' is not in Active/Approved state.") % comp.name})
                        continue
                    comp_cache[raw_comp.lower()] = comp
                else:
                    errors.append({'row': row_idx, 'field': 'Competency', 'message': _("Competency '%s' was not found.") % raw_comp})
                    continue

            # 4. Resolve Required Level
            lvl_digit = None
            lvl_match = re.search(r'[1-4]', raw_lvl)
            if lvl_match:
                lvl_digit = lvl_match.group(0)
            if not lvl_digit or lvl_digit not in ('1', '2', '3', '4'):
                errors.append({'row': row_idx, 'field': 'Required Level', 'message': _("Invalid Required Level '%s'. Must be 1, 2, 3, or 4.") % raw_lvl})
                continue

            # Check for duplicate competency within the same job and operating unit in this file
            dedup_key = (job.id, ou.id if ou else False, comp.id)
            if dedup_key in seen_job_comp:
                first_row = seen_job_comp[dedup_key]
                ou_str = f" and Operating Unit '{ou.name}'" if ou else ""
                errors.append({
                    'row': row_idx,
                    'field': 'Competency',
                    'message': _("Duplicate entry for Competency '%(comp)s' under Job Position '%(job)s'%(ou)s (first defined on row %(first_row)s).") % {
                        'comp': comp.name,
                        'job': job.name,
                        'ou': ou_str,
                        'first_row': first_row,
                    }
                })
                continue
            seen_job_comp[dedup_key] = row_idx

            # 5. Resolve Weight
            parsed_weight = 1.0
            if raw_weight not in (None, ''):
                try:
                    parsed_weight = float(raw_weight)
                    if parsed_weight <= 0:
                        parsed_weight = 1.0
                except (ValueError, TypeError):
                    parsed_weight = 1.0

            # 6. Resolve Override Default
            override_bool = False
            if isinstance(raw_override, bool):
                override_bool = raw_override
            elif str(raw_override).strip().lower() in ('true', 'yes', '1', 'y', 't'):
                override_bool = True

            # 7. Resolve Row Status
            row_status = 'draft'
            if self.is_admin:
                if self.import_state == 'approved':
                    row_status = 'approved'
                elif self.import_state == 'from_excel':
                    if raw_status.strip().lower() in ('approved', 'approve', 'app'):
                        row_status = 'approved'
                    else:
                        row_status = 'draft'
                else:
                    row_status = 'draft'
            else:
                row_status = 'draft'

            valid_rows.append({
                'row_idx': row_idx,
                'job': job,
                'ou': ou,
                'competency': comp,
                'level': lvl_digit,
                'weight': parsed_weight,
                'override': override_bool,
                'status': row_status,
            })

        # If blocking validation errors were found, halt and present them clearly
        if errors:
            err_html = """
            <div class="alert alert-danger" role="alert" style="margin-bottom: 15px;">
                <h4 style="margin-top: 0;"><i class="fa fa-exclamation-triangle"></i> Excel Validation Failed</h4>
                <p>Found <strong>%d error(s)</strong> across <strong>%d rows read</strong>. No records were modified. Please correct the Excel sheet and upload again.</p>
            </div>
            <table class="table table-bordered table-striped" style="font-size: 13px;">
                <thead>
                    <tr style="background-color: #541718; color: #FFFFFF;">
                        <th style="width: 80px; text-align: center;">Row #</th>
                        <th style="width: 160px;">Column Field</th>
                        <th>Error Description</th>
                    </tr>
                </thead>
                <tbody>
            """ % (len(errors), total_rows_read)

            for err in errors[:50]:
                err_html += f"""
                    <tr>
                        <td style="text-align: center; font-weight: bold;">{err['row']}</td>
                        <td><span class="badge bg-secondary">{err['field']}</span></td>
                        <td style="color: #c92a2a;">{err['message']}</td>
                    </tr>
                """
            if len(errors) > 50:
                err_html += f"""
                    <tr>
                        <td colspan="3" class="text-center text-muted">... and {len(errors) - 50} more errors.</td>
                    </tr>
                """
            err_html += "</tbody></table>"

            self.result_summary = err_html
            self.state = 'done'
            return {
                'type': 'ir.actions.act_window',
                'res_model': 'competency.role.mapping.import.wizard',
                'res_id': self.id,
                'view_mode': 'form',
                'target': 'new',
            }

        if not valid_rows:
            raise UserError(_("No valid data rows found in the uploaded file."))

        # Group valid rows by Mapping Key: (job_id, ou_id)
        grouped_mappings = {}
        for r in valid_rows:
            ou_id = r['ou'].id if r['ou'] else False
            m_key = (r['job'].id, ou_id)
            grouped_mappings.setdefault(m_key, {
                'job': r['job'],
                'ou': r['ou'],
                'status': r['status'],
                'lines': [],
            })
            grouped_mappings[m_key]['lines'].append(r)
            # If any line requested approval and user is admin, promote group status
            if r['status'] == 'approved' and self.is_admin:
                grouped_mappings[m_key]['status'] = 'approved'

        # Check for duplicate competencies within the same mapping group
        for m_key, group_data in grouped_mappings.items():
            seen_comps = {}
            for l in group_data['lines']:
                c_id = l['competency'].id
                if c_id in seen_comps:
                    raise UserError(_(
                        "Duplicate competency detected in Job Position '%s': '%s' was entered on both row %d and row %d."
                    ) % (group_data['job'].name, l['competency'].name, seen_comps[c_id], l['row_idx']))
                seen_comps[c_id] = l['row_idx']

        # Process role mappings
        Mapping = self.env['competency.role.mapping'].with_context(force_write=True, skip_mapping_unique_check=True)
        MappingLine = self.env['competency.role.mapping.line'].with_context(force_write=True)

        processed_mappings = self.env['competency.role.mapping']
        log_rows = []

        for (job_id, ou_id), group_data in grouped_mappings.items():
            job = group_data['job']
            ou = group_data['ou']
            target_status = group_data['status']

            # Lookup existing mappings for this Job Position + OU
            domain = [('job_position_id', '=', job_id)]
            if ou_id:
                domain.extend([('is_operating_unit_specific', '=', True), ('operating_unit_ids', 'in', [ou_id])])
            else:
                domain.append(('is_operating_unit_specific', '=', False))

            existing_mappings = Mapping.search(domain)
            existing_approved = existing_mappings.filtered(lambda m: m.state == 'approved')
            existing_draft = existing_mappings.filtered(lambda m: m.state in ('draft', 'under_approval'))

            action_taken = ""
            active_rec = False

            if self.import_mode == 'override':
                # Pick existing approved or draft mapping to update in place
                target_mapping = existing_approved[0] if existing_approved else (existing_draft[0] if existing_draft else False)
                if target_mapping:
                    # Clear existing lines and replace
                    target_mapping.line_ids.unlink()
                    for line_info in group_data['lines']:
                        MappingLine.create({
                            'mapping_id': target_mapping.id,
                            'competency_id': line_info['competency'].id,
                            'required_proficiency': line_info['level'],
                            'weight': line_info['weight'],
                            'override_default': line_info['override'],
                        })
                    
                    update_vals = {
                        'change_description': _("Overridden via Excel Import ('%s').") % (self.filename or 'Excel file'),
                    }
                    if target_status == 'approved' and self.is_admin:
                        update_vals.update({
                            'state': 'approved',
                            'approved_by_id': self.env.user.id,
                            'approval_date': fields.Datetime.now(),
                        })
                    elif target_status == 'draft':
                        update_vals.update({'state': 'draft'})

                    target_mapping.write(update_vals)
                    target_mapping.message_post(body=_("Competency mapping lines updated via Excel Import."))
                    active_rec = target_mapping
                    action_taken = _("Overridden in place (Version %s)") % target_mapping.version
                else:
                    # No existing mapping exists -> Create initial version
                    create_vals = {
                        'job_position_id': job_id,
                        'version': 'v1.0',
                        'state': target_status if self.is_admin else 'draft',
                        'change_description': _("Created via Excel Import ('%s').") % (self.filename or 'Excel file'),
                        'is_operating_unit_specific': bool(ou_id),
                        'operating_unit_ids': [(6, 0, [ou_id])] if ou_id else False,
                    }
                    if target_status == 'approved' and self.is_admin:
                        create_vals.update({
                            'approved_by_id': self.env.user.id,
                            'approval_date': fields.Datetime.now(),
                        })
                    active_rec = Mapping.create(create_vals)
                    for line_info in group_data['lines']:
                        MappingLine.create({
                            'mapping_id': active_rec.id,
                            'competency_id': line_info['competency'].id,
                            'required_proficiency': line_info['level'],
                            'weight': line_info['weight'],
                            'override_default': line_info['override'],
                        })
                    action_taken = _("Created Initial Version (v1.0)")

            else:
                # Mode: 'new_version'
                # Determine new version string
                if existing_approved:
                    highest_version = existing_approved[0].version
                    new_version = Mapping._bump_version(highest_version)
                    # Archive existing approved mapping
                    existing_approved.write({'state': 'archived'})
                    for p in existing_approved:
                        p.message_post(body=_("Archived and superseded by imported version %s.") % new_version)
                else:
                    new_version = 'v1.0'

                # Also archive any old draft mappings to avoid conflicts
                if existing_draft:
                    existing_draft.write({'state': 'archived'})

                create_vals = {
                    'job_position_id': job_id,
                    'version': new_version,
                    'state': target_status if self.is_admin else 'draft',
                    'change_description': _("Imported version created from Excel file '%s'.") % (self.filename or 'Excel file'),
                    'is_operating_unit_specific': bool(ou_id),
                    'operating_unit_ids': [(6, 0, [ou_id])] if ou_id else False,
                }
                if target_status == 'approved' and self.is_admin:
                    create_vals.update({
                        'approved_by_id': self.env.user.id,
                        'approval_date': fields.Datetime.now(),
                    })
                active_rec = Mapping.create(create_vals)
                for line_info in group_data['lines']:
                    MappingLine.create({
                        'mapping_id': active_rec.id,
                        'competency_id': line_info['competency'].id,
                        'required_proficiency': line_info['level'],
                        'weight': line_info['weight'],
                        'override_default': line_info['override'],
                    })
                action_taken = _("New Version Created (%s)") % new_version

            if active_rec:
                if active_rec.state == 'approved':
                    active_rec._notify_assigned_employees()
                processed_mappings |= active_rec
                log_rows.append({
                    'job_name': job.name,
                    'ou_name': ou.name if ou else _("Bank-Wide (Global)"),
                    'action': action_taken,
                    'comp_count': len(group_data['lines']),
                    'status': dict(active_rec._fields['state'].selection).get(active_rec.state, active_rec.state),
                    'version': active_rec.version,
                })

        # Build Success Summary Report
        summary_html = f"""
        <div class="alert alert-success" role="alert" style="margin-bottom: 20px;">
            <h4 style="margin-top: 0;"><i class="fa fa-check-circle"></i> Excel Import Completed Successfully</h4>
            <p>Processed <strong>{total_rows_read} Excel rows</strong> across <strong>{len(grouped_mappings)} Job Positions</strong>. Configured <strong>{len(valid_rows)} Competency Requirements</strong>.</p>
        </div>
        <table class="table table-bordered table-hover" style="font-size: 13.5px;">
            <thead>
                <tr style="background-color: #541718; color: #FFFFFF;">
                    <th>Job Position</th>
                    <th>Operating Unit Scope</th>
                    <th style="text-align: center;">Version</th>
                    <th style="text-align: center;">Competencies</th>
                    <th>Action Taken</th>
                    <th style="text-align: center;">Final Status</th>
                </tr>
            </thead>
            <tbody>
        """
        for l in log_rows:
            status_badge = 'bg-success' if l['status'] == 'Approved' else 'bg-warning text-dark'
            summary_html += f"""
                <tr>
                    <td style="font-weight: 600;">{l['job_name']}</td>
                    <td><span class="text-muted">{l['ou_name']}</span></td>
                    <td style="text-align: center;"><strong>{l['version']}</strong></td>
                    <td style="text-align: center;"><span class="badge bg-info text-dark">{l['comp_count']} lines</span></td>
                    <td>{l['action']}</td>
                    <td style="text-align: center;"><span class="badge {status_badge}">{l['status']}</span></td>
                </tr>
            """
        summary_html += "</tbody></table>"

        self.write({
            'result_summary': summary_html,
            'state': 'done',
            'created_mapping_ids': [(6, 0, processed_mappings.ids)],
        })

        return {
            'type': 'ir.actions.act_window',
            'res_model': 'competency.role.mapping.import.wizard',
            'res_id': self.id,
            'view_mode': 'form',
            'target': 'new',
        }

    def action_view_processed_mappings(self):
        """Opens list view of all created or updated role mappings."""
        self.ensure_one()
        return {
            'name': _('Processed Role-Competency Mappings'),
            'type': 'ir.actions.act_window',
            'res_model': 'competency.role.mapping',
            'view_mode': 'list,form',
            'domain': [('id', 'in', self.created_mapping_ids.ids)],
            'target': 'current',
        }
